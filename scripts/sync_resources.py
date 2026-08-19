#!/usr/bin/env python3
"""Download Google Sheet resource DB and write data/resources.json."""

from __future__ import annotations

import json
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Missing dependency: openpyxl. Install with: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

try:
    import urllib.request
except ImportError:
    urllib = None  # type: ignore

SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1q6_dwNBk-uhuj0-m9iiQvEvHpei-kUOo6hUVl5YzH0k/export?format=xlsx"
)

ROOT = Path(__file__).resolve().parent.parent
LOCAL_FALLBACK = ROOT / "resource_database_ready_for_google_sheets.xlsx"
OUT_PATH = ROOT / "data" / "resources.json"
JS_OUT_PATH = ROOT / "js" / "resources-data.js"

SHEET_NAMES = ("States", "Programs", "Resource_Types", "Organizations", "Resources")


def is_truthy(value) -> bool:
    if value is True:
        return True
    if value is False or value is None:
        return False
    s = str(value).strip().lower()
    return s in {"true", "yes", "1"}


def cell_str(value) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return s if s else None


def sheet_to_dicts(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        item = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            item[h] = row[i] if i < len(row) else None
        out.append(item)
    return out


def normalize_lookup_rows(rows: list[dict]) -> list[dict]:
    """Convert cells to JSON-friendly values (keep keys as-is)."""
    normalized = []
    for row in rows:
        item = {}
        for k, v in row.items():
            if isinstance(v, datetime):
                item[k] = v.isoformat()
            elif v is None:
                item[k] = None
            elif isinstance(v, (int, float, bool)):
                item[k] = v
            else:
                item[k] = str(v).strip() if str(v).strip() != "" else None
        normalized.append(item)
    return normalized


def normalize_resource(row: dict) -> dict:
    item = {}
    for k, v in row.items():
        if isinstance(v, datetime):
            item[k] = v.isoformat()
        else:
            item[k] = v

    item["active"] = is_truthy(item.get("active"))

    rtid = item.get("resource_type_id")
    if rtid is None or (isinstance(rtid, str) and not rtid.strip()):
        item["resource_type_id"] = None
    else:
        # keep as string (even if numeric in sheet)
        if isinstance(rtid, float) and rtid == int(rtid):
            item["resource_type_id"] = str(int(rtid))
        else:
            item["resource_type_id"] = str(rtid).strip()

    for key in ("state_id", "program_id"):
        val = item.get(key)
        if val is None or (isinstance(val, str) and not str(val).strip()):
            item[key] = None
        else:
            item[key] = str(val).strip().upper()

    phone = cell_str(item.get("phone"))
    item["phone"] = phone

    url = cell_str(item.get("url"))
    item["url"] = url

    # stringify other non-null scalars for consistency where useful
    for k, v in list(item.items()):
        if k in ("active", "phone", "url", "resource_type_id", "state_id", "program_id"):
            continue
        if v is None:
            continue
        if isinstance(v, bool):
            continue
        if isinstance(v, (int, float)):
            # leave numbers as-is unless they look like ids that should be strings
            if k.endswith("_id") or k in ("id", "organization_id"):
                if isinstance(v, float) and v == int(v):
                    item[k] = str(int(v))
                else:
                    item[k] = str(v)
            continue
        s = str(v).strip()
        item[k] = s if s else None

    return item


def download_xlsx(dest: Path) -> bool:
    try:
        req = urllib.request.Request(
            SHEET_URL,
            headers={"User-Agent": "Mozilla/5.0 (compatible; SaludSimpleSync/1.0)"},
        )
        with urllib.request.urlopen(req, timeout=60) as resp:
            data = resp.read()
        if len(data) < 100:
            print("Download too small; treating as failure.", file=sys.stderr)
            return False
        # xlsx is a zip; should start with PK
        if not data.startswith(b"PK"):
            print("Download did not look like xlsx; treating as failure.", file=sys.stderr)
            return False
        dest.write_bytes(data)
        print(f"Downloaded sheet to {dest}")
        return True
    except Exception as e:
        print(f"Download failed: {e}", file=sys.stderr)
        return False


def load_workbook_path() -> Path:
    tmp = Path(tempfile.gettempdir()) / "ellie_resources_sheet.xlsx"
    if download_xlsx(tmp):
        return tmp
    if LOCAL_FALLBACK.exists():
        print(f"Falling back to local file: {LOCAL_FALLBACK}")
        return LOCAL_FALLBACK
    print("No spreadsheet available (download failed and local fallback missing).", file=sys.stderr)
    sys.exit(1)


def main() -> None:
    xlsx_path = load_workbook_path()
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)

    missing = [n for n in SHEET_NAMES if n not in wb.sheetnames]
    if missing:
        print(f"Missing sheets: {missing}. Found: {wb.sheetnames}", file=sys.stderr)
        sys.exit(1)

    states = normalize_lookup_rows(sheet_to_dicts(wb["States"]))
    programs = normalize_lookup_rows(sheet_to_dicts(wb["Programs"]))
    resource_types = normalize_lookup_rows(sheet_to_dicts(wb["Resource_Types"]))
    organizations = normalize_lookup_rows(sheet_to_dicts(wb["Organizations"]))
    resources_raw = sheet_to_dicts(wb["Resources"])
    wb.close()

    resources = []
    for row in resources_raw:
        norm = normalize_resource(row)
        if norm.get("active") is True:
            resources.append(norm)

    payload = {
        "updated": datetime.now(timezone.utc).isoformat(),
        "states": states,
        "programs": programs,
        "resource_types": resource_types,
        "organizations": organizations,
        "resources": resources,
    }

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    json_text = json.dumps(payload, ensure_ascii=False, indent=2)
    OUT_PATH.write_text(json_text, encoding="utf-8")
    print(f"Wrote {OUT_PATH}")

    JS_OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    # Compact JSON for browser embed (file:// friendly; no fetch/CORS)
    js_payload = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    JS_OUT_PATH.write_text(
        "window.SALUD_RESOURCES = " + js_payload + ";\n",
        encoding="utf-8",
    )
    print(f"Wrote {JS_OUT_PATH}")
    print(
        f"counts: states={len(states)} programs={len(programs)} "
        f"resource_types={len(resource_types)} organizations={len(organizations)} "
        f"active_resources={len(resources)}"
    )


if __name__ == "__main__":
    main()
